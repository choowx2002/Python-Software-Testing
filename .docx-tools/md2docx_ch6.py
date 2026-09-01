#!/usr/bin/env python3
"""Convert chapter6-draft markdown into a Word docx (headings, tables, red placeholders)."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "docs/chapter6-draft-2026-08-29.md"
OUT = "docs/chapter6-draft-2026-08-29.docx"
RED = RGBColor(0xC0, 0x00, 0x00)
HDR_FILL = "D9D9D9"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

def add_runs(p, text, size=11):
    """Parse **bold** and `code` inline markers."""
    pos = 0
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.font.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(size - 0.5)
        else:
            r = p.add_run(tok)
        r.font.size = Pt(size)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def add_table(rows):
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.paragraphs[0].text = ""
            p = c.paragraphs[0]
            if ri == 0:
                r = p.add_run(cell); r.font.bold = True; r.font.size = Pt(9.5)
                shade(c, HDR_FILL)
            else:
                add_runs(p, cell, size=9.5)
    # fixed layout
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    total = len(rows[0])
    width = 15.9 / total
    for row in t.rows:
        for cell in row.cells:
            cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue
    # tables
    if line.startswith("|"):
        tbl = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):  # skip separator row
                tbl.append(cells)
            i += 1
        add_table(tbl)
        continue
    # headings
    m = re.match(r"^(#{1,4})\s+(.*)$", line)
    if m:
        level = min(len(m.group(1)), 3)
        text = m.group(2).replace("**", "")
        doc.add_heading(text, level=level)
        i += 1
        continue
    # paragraph
    p = doc.add_paragraph()
    if "【待填" in line:
        r = p.add_run(line); r.font.color.rgb = RED
    else:
        add_runs(p, line)
    p.paragraph_format.space_after = Pt(6)
    i += 1

doc.core_properties.title = "Chapter 6 Draft — Implementation, Testing and Evaluation"
doc.save(OUT)
print("saved:", OUT)
