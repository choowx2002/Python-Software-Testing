#!/usr/bin/env python3
"""Convert the UAT test plan markdown into a Word docx (headings, tables, lists, blockquotes)."""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else "docs/UAT-test-plan.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "docs/UAT-test-plan.docx"
HDR_FILL = "D9D9D9"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")

def add_runs(p, text, size=11):
    """Parse **bold** and `code` inline markers."""
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.font.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(size - 0.5)
        else:
            r = p.add_run(tok)
        r.font.size = Pt(size)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def col_widths(headers):
    """Return proportional column widths (cm) keyed on the table shape."""
    n = len(headers)
    first = headers[0] if headers else ""
    if first == "Case ID":  # UAT case tables (8 columns)
        return [1.2, 1.7, 1.0, 2.1, 3.8, 3.8, 1.2, 1.1]
    if n == 2:              # metadata / reference / environment tables
        return [4.5, 11.4]
    if n == 7:              # defect log
        return [1.6, 2.2, 3.6, 1.6, 1.8, 1.6, 3.5]
    if n == 4:              # sign-off
        return [3.0, 4.0, 4.4, 4.5]
    if n == 5:              # test case summary
        return [1.6, 4.6, 3.2, 3.0, 3.5]
    w = 15.9 / n            # fallback: equal widths
    return [w] * n

def add_table(rows):
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    widths = col_widths(rows[0])
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.rows[ri].cells[ci]
            c.paragraphs[0].text = ""
            p = c.paragraphs[0]
            if ri == 0:
                r = p.add_run(cell); r.font.bold = True; r.font.size = Pt(9)
                shade(c, HDR_FILL)
            else:
                add_runs(p, cell, size=9)
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if ci < len(widths):
                cell.width = Cm(widths[ci])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip():
        i += 1
        continue
    # tables
    if line.startswith("|"):
        tbl = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):  # skip separator row
                tbl.append(cells)
            i += 1
        add_table(tbl)
        continue
    # horizontal rules
    if re.match(r"^-{3,}$", line.strip()):
        i += 1
        continue
    # headings
    m = re.match(r"^(#{1,4})\s+(.*)$", line)
    if m:
        level = min(len(m.group(1)), 3)
        doc.add_heading(m.group(2).replace("**", ""), level=level)
        i += 1
        continue
    # blockquote
    if line.startswith(">"):
        p = doc.add_paragraph()
        r = p.add_run(line.lstrip("> ").strip())
        r.font.italic = True
        r.font.size = Pt(10)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(6)
        i += 1
        continue
    # ordered list items: "1. text" or "1) text"
    m = re.match(r"^(\d+)[.)]\s+(.*)$", line)
    if m:
        p = doc.add_paragraph()
        add_runs(p, f"{m.group(1)}.  {m.group(2)}", size=11)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(3)
        i += 1
        continue
    # plain paragraph
    p = doc.add_paragraph()
    add_runs(p, line)
    p.paragraph_format.space_after = Pt(6)
    i += 1

doc.core_properties.title = "Testmate — User Acceptance Test Plan (Functional Requirements)"
doc.save(OUT)
print("saved:", OUT)
